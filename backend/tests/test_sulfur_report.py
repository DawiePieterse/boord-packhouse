"""Batch number, size and second source on a run, the dose a tent starts with
next time, the size list, and the daily report with its Excel and PDF."""
import io
import json
from datetime import datetime, timedelta, timezone

from tests.test_sulfur import run, sync


def _local_noon():
    return datetime.now().astimezone().replace(hour=12, minute=0, second=0, microsecond=0)


def test_batch_size_and_second_source_are_kept(client):
    more = json.dumps([{"block_id": "8a", "block_name": "Block 8a", "variety": "Third Month Red"}])
    sync(client, run(batch_no="T1-4104-1057-B15/8a-K7", size="XL", more_sources=more))
    r = client.get("/api/sulfur/runs").json()["runs"][0]
    assert (r["batch_no"], r["size"]) == ("T1-4104-1057-B15/8a-K7", "XL")
    assert json.loads(r["more_sources"])[0]["block_id"] == "8a"


def test_a_malformed_second_source_is_refused(client):
    out = sync(client, run(more_sources="not json"), run("r2", more_sources=json.dumps([{"block_id": "1", "supplier_id": 2}])))
    assert out["synced"] == [] and len(out["rejected"]) == 2


def test_a_tent_starts_with_the_dose_last_used(client):
    assert client.get("/api/tents").json()[0]["last_grams"] is None
    sync(client, run("a", started=datetime.now(timezone.utc) - timedelta(hours=2), grams=400))
    sync(client, run("b", grams=410))
    assert client.get("/api/tents").json()[0]["last_grams"] == 410
    assert client.get("/api/tents").json()[1]["last_grams"] is None


def test_the_size_list_is_edited_in_settings(admin_client):
    assert json.loads(admin_client.get("/api/sulfur/settings").json()["sizes"]) == ["L", "XL", "XXL", "Mixed"]
    body = {k: v for k, v in admin_client.get("/api/settings/sulfur").json()["settings"].items()
            if k not in ("id", "sizes", "variety_factors")}
    assert admin_client.put("/api/settings/sulfur", json={**body, "sizes": [" M ", "L", "M", ""]}).status_code == 200
    assert json.loads(admin_client.get("/api/sulfur/settings").json()["sizes"]) == ["M", "L"]
    assert admin_client.put("/api/settings/sulfur", json=body).status_code == 200  # no list sent: unchanged
    assert json.loads(admin_client.get("/api/sulfur/settings").json()["sizes"]) == ["M", "L"]
    assert admin_client.put("/api/settings/sulfur", json={**body, "sizes": [" "]}).status_code == 422


def _two_days(client):
    noon = _local_noon()
    for uuid, days_ago, tent, grams, crates in (("a", 1, 1, 425, 10), ("b", 0, 1, 400, 24), ("c", 0, 2, 400, 12)):
        sync(client, run(uuid, started=noon - timedelta(days=days_ago), tent_id=tent, tent_name=f"Tent {tent}",
                         grams=grams, crates=crates, size="Mixed", batch_no=f"B-{uuid}"))


def test_the_report_needs_the_pin(client):
    assert client.get("/api/settings/sulfur/report?from=2026-10-01").status_code == 401


def test_daily_totals_per_tent_and_day(admin_client):
    _two_days(admin_client)
    day = _local_noon().date()
    data = admin_client.get(f"/api/settings/sulfur/report?from={day - timedelta(days=1)}&to={day}").json()
    assert [r["batch_no"] for r in data["runs"]] == ["B-a", "B-b", "B-c"]
    today = {d["tent"]: d for d in data["days"] if d["date"] == day.isoformat()}
    assert today["All tents"]["grams"] == 800 and today["All tents"]["lugs"] == 36 and today["All tents"]["runs"] == 2
    assert today["Tent 1"]["grams"] == 400
    only = admin_client.get(f"/api/settings/sulfur/report?from={day}&tent_id=2").json()
    assert [r["batch_no"] for r in only["runs"]] == ["B-c"]


def test_excel_and_pdf_exports(admin_client):
    from openpyxl import load_workbook
    _two_days(admin_client)
    day = _local_noon().date()
    q = f"from={day - timedelta(days=1)}&to={day}"
    x = admin_client.get(f"/api/settings/sulfur/report/xlsx?{q}")
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Sulfur room log", "Usage per day"]
    assert wb["Sulfur room log"]["B2"].value == "B-a" and wb["Usage per day"].max_row == 1 + 5
    p = admin_client.get(f"/api/settings/sulfur/report/pdf?{q}")
    assert p.status_code == 200 and p.content.startswith(b"%PDF")
    assert admin_client.get("/api/settings/sulfur/report/pdf?from=2026-10-05&to=2026-10-01").status_code == 422
