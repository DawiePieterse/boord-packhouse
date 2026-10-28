"""Delivery notes: the pallets' editable sync, a number unique among live
notes, the Pallet Stack pallets on it (a copy of each, each on one note), the
settings it picks from, and the paper note as a PDF."""
from datetime import datetime, timezone

import pytest

from tests.test_consignment import cons, pallet
from tests.test_consignment import sync as sync_pallets

NOW = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def note(uuid="n1", note_no=1335, pallets=("p1",), **kw):
    return {"uuid": uuid, "note_no": note_no, "created_at": NOW, "receiver_id": 2,
            "receiver": "Boeretrust\nTshwane Mark\nPretoria", "transporter_id": 2, "transporter": "Go Fresh",
            "loaded_on": "2026-10-26", "offloaded_on": "2026-10-27", "pallets": [{"uuid": u} for u in pallets], **kw}


def sync(client, *notes):
    r = client.post("/api/deliverynote/sync", json={"device": "phone-1", "notes": list(notes)})
    assert r.status_code == 200
    return r.json()


def notes(client):
    return client.get("/api/deliverynote/state").json()["notes"]


@pytest.fixture
def pallets(client):
    sync_pallets(client, [cons()], [pallet("p1"), pallet("p2", barcode="B2")])


def test_sync_is_idempotent_and_editable_and_keeps_who_made_it(client, pallets):
    assert sync(client, note(made_by="W1", made_by_name="Thandi"))["synced"] == ["n1"]
    sync(client, note(seal_no="S9", made_by="W2", made_by_name="Sipho", changed_by_name="Sipho"))
    [n] = notes(client)
    assert (n["seal_no"], n["made_by_name"], n["changed_by_name"], n["receiver_id"]) == ("S9", "Thandi", "Sipho", 2)


def test_the_note_keeps_its_own_copy_of_each_pallet(client, pallets):
    sync(client, note(pallets=("p1", "p2")))
    sync_pallets(client, pallets=[pallet("p1", supplier_name="Changed later")])
    [p1, p2] = notes(client)[0]["pallets"]
    assert (p1["pallet_ref"], p1["supplier_name"], p1["lines"][0]["cartons"]) == ("001", "Own fruit", 60)
    assert (p2["uuid"], p2["barcode"], p2["market_name"]) == ("p2", "B2", "USA Sea Freight")


def test_a_pallet_is_on_one_live_note(client, pallets):
    sync(client, note())
    assert "already on note 1335" in sync(client, note("n2", 1336))["rejected"][0]["error"]
    sync(client, note(removed=True))
    assert sync(client, note("n2", 1336))["synced"] == ["n2"]


def test_number_is_unique_until_the_note_is_removed(client, pallets):
    sync(client, note())
    assert "already used" in sync(client, note("n2", pallets=("p2",)))["rejected"][0]["error"]
    sync(client, note(removed=True))
    assert sync(client, note("n2", pallets=("p2",)))["synced"] == ["n2"]


@pytest.mark.parametrize("bad", [
    {"note_no": 0}, {"receiver": " "}, {"loaded_on": ""}, {"offloaded_on": "27/10"}, {"pallets": ()},
    {"pallets": ("nope",)}, {"pallets": ("p1", "p1")},
])
def test_bad_notes_are_rejected_not_stored(client, pallets, bad):
    assert sync(client, note(**bad))["rejected"][0]["uuid"] == "n1"
    assert notes(client) == []


def test_pdf_is_the_paper_note_for_local_notes_only(client, pallets):
    sync_pallets(client, pallets=[pallet("p3", barcode="B3", market_id=5, market_name="Local")])
    sync(client, note(pallets=("p3",), cooled_to="8"), note("n2", 1336, pallets=("p1", "p2")))
    r = client.get("/api/deliverynote/n1/pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert "afleweringsnota_B-1335.pdf" in r.headers["content-disposition"]
    assert client.get("/api/deliverynote/n2/pdf").status_code == 409  # USA pallets: their papers instead
    assert client.get("/api/deliverynote/nope/pdf").status_code == 404


def test_options_come_from_settings_which_need_the_pin(client, admin_client):
    o = client.get("/api/deliverynote/options").json()
    assert ([t["name"] for t in o["transporters"]], o["cooled_to"], o["last_no"]) == (["Go Fresh", "Self"], 8.0, None)  # A to Z
    assert [r["name"] for r in o["receivers"]] == ["Boeretrust", "Laughing Waters Lietsjies"]
    assert (o["receivers"][0]["market"], o["receivers"][0]["city"]) == ("Tshwane Mark", "Pretoria")
    r = admin_client.post("/api/settings/receivers", json={"name": " Dapper ", "market": "Joburg Mark", "city": "Johannesburg"})
    assert r.json()["name"] == "Dapper"
    admin_client.post("/api/settings/receivers", json={"name": "agri Co"})  # capitals or not, A to Z
    assert [r["name"] for r in client.get("/api/deliverynote/options").json()["receivers"]] == [
        "agri Co", "Boeretrust", "Dapper", "Laughing Waters Lietsjies"]
    admin_client.put("/api/settings/transporters/1", json={"name": "Self", "active": False})
    assert admin_client.put("/api/settings/deliverynote", json={"cooled_to": 5}).json()["cooled_to"] == 5
    t = admin_client.post("/api/settings/targetreceivers", json={"name": " Seasons Farm Fresh ", "address": " Philadelphia ", "market": "usa"}).json()
    assert (t["name"], t["address"], t["market"]) == ("Seasons Farm Fresh", "Philadelphia", "USA")
    assert admin_client.post("/api/settings/targetreceivers", json={"name": "X", "market": "UK"}).status_code == 422
    admin_client.post("/api/settings/targetreceivers", json={"name": "Old one", "market": "EU", "active": False})
    assert [r["name"] for r in client.get("/api/deliverynote/options").json()["target_receivers"]] == ["Seasons Farm Fresh"]
    assert admin_client.put("/api/settings/deliverynote", json={"cooled_to": 99}).status_code == 422
    o = client.get("/api/deliverynote/options").json()
    assert ([t["name"] for t in o["transporters"]], len(o["receivers"]), o["cooled_to"]) == (["Go Fresh"], 4, 5)
    admin_client.post("/api/admin/logout")
    assert client.get("/api/settings/receivers").status_code == 401
    assert client.put("/api/settings/deliverynote", json={"cooled_to": 5}).status_code == 401


def test_delete_only_what_no_live_note_uses(client, admin_client, pallets):
    sync(client, note())  # receiver 2, transporter 2
    assert admin_client.delete("/api/settings/receivers/2").status_code == 409
    assert admin_client.delete("/api/settings/transporters/2").status_code == 409
    assert admin_client.delete("/api/settings/transporters/1").status_code == 200
    sync(client, note(removed=True))
    assert admin_client.delete("/api/settings/receivers/2").status_code == 200


def test_last_number_follows_notes_and_settings_can_change_it(client, admin_client, pallets):
    last = lambda: client.get("/api/deliverynote/options").json()["last_no"]  # noqa: E731
    sync(client, note())
    assert last() == 1335
    # Mistyped 13350 then corrected: Settings sets it back, a lower note never pulls it down.
    assert admin_client.put("/api/settings/deliverynote", json={"cooled_to": 8, "last_no": 1340}).status_code == 200
    assert last() == 1340
    sync(client, note("n2", 1336, pallets=("p2",)))
    assert last() == 1340
    sync(client, note("n2", 1341, pallets=("p2",)))
    assert last() == 1341
    assert admin_client.put("/api/settings/deliverynote", json={"cooled_to": 8, "last_no": 0}).status_code == 422
    admin_client.put("/api/settings/deliverynote", json={"cooled_to": 8, "last_no": None})
    assert last() == 1341  # cleared: the highest note ever used


def test_usa_forms_are_drawn_from_the_note(client, pallets):
    from io import BytesIO

    from openpyxl import load_workbook
    sync(client, note(pallets=("p1", "p2")))
    r = client.get("/api/deliverynote/n1/usa_forms.xlsx")
    assert r.status_code == 200 and 'usa_forms_LW-1335.xlsx' in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Phyto", "Consignment", "Inspection Sheet"]
    p, c, ins = wb["Phyto"], wb["Consignment"], wb["Inspection Sheet"]
    assert (p["D5"].value, p["B19"].value, p["J20"].value, p["F16"].value, p["B21"].value) == (
        "LW-1335", "060099005963004759", "002", 120, None)
    assert (c["J4"].value, c["C15"].value, c["E4"].value, c["K16"].value) == ("LW-1335", "B2", "Boeretrust", 120)
    assert client.get("/api/deliverynote/nope/usa_forms.xlsx").status_code == 404
    assert client.get("/api/deliverynote/n1/usa.xlsx").status_code == 404  # no simple sheets
    assert (ins["A11"].value, ins["D11"].value) == ("B2", "USA")


def test_export_forms_also_come_as_one_pdf(client, pallets):
    sync(client, note(pallets=("p1", "p2")))
    r = client.get("/api/deliverynote/n1/usa_forms.pdf")
    assert r.status_code == 200 and 'usa_forms_LW-1335.pdf' in r.headers["content-disposition"]
    assert r.content.startswith(b"%PDF") and r.content.count(b"/Type /Page\n") == 3  # a page a sheet
    assert client.get("/api/deliverynote/n1/eu_forms.pdf").content.count(b"/Type /Page\n") == 2  # no phyto
    assert client.get("/api/deliverynote/n1/uk_forms.pdf").status_code == 404
    assert client.get("/api/deliverynote/nope/usa_forms.pdf").status_code == 404


def test_usa_forms_keep_phyto_on_one_page_and_print_the_target_receiver():
    import json as _json
    from io import BytesIO

    from openpyxl import load_workbook

    import export_forms
    from models import DeliveryNote
    line = {"cartons": 192, "packaging_name": "Rhino 4.5 kg XL", "carton_kg": 4.5}
    n = DeliveryNote(uuid="x", note_no=7, created_at=datetime.now(), loaded_on="2026-10-26", receiver="AgriLink",
                     pallets=_json.dumps([{"pallet_ref": str(k), "barcode": "", "date_code": "", "lines": [line]} for k in range(27)]))
    wb = load_workbook(BytesIO(export_forms.forms(n, "Seasons Farm Fresh\nToronto")))
    p = wb["Phyto"]
    assert wb.sheetnames == ["Phyto", "Consignment", "Inspection Sheet"]
    assert (p["A45"].value, p.page_setup.fitToHeight) == (27, 1)
    c = wb["Consignment"]
    assert (c["J8"].value, c["A8"].value, c["C10"].value, c["A11"].value) == (
        "Seasons Farm Fresh\nToronto", "PUC code", "4050373828202", "GlobalG.A.P. Certified")


def test_export_papers_print_the_notes_own_target_receiver(client, admin_client, pallets):
    from io import BytesIO

    from openpyxl import load_workbook
    t = admin_client.post("/api/settings/targetreceivers", json={"name": "Rotterdam Fresh", "market": "EU"}).json()
    sync(client, note(pallets=("p1", "p2"), target_receiver_id=t["id"], target_receiver="Rotterdam Fresh\nRotterdam"))
    wb = load_workbook(BytesIO(client.get("/api/deliverynote/n1/eu_forms.xlsx").content))
    assert wb.sheetnames == ["Consignment", "Inspection Sheet"]
    c, ins = wb["Consignment"], wb["Inspection Sheet"]
    assert (c["J7"].value, c["J8"].value, ins["D10"].value) == ("EU", "Rotterdam Fresh\nRotterdam", "EU")
    assert client.get("/api/deliverynote/n1/uk_forms.xlsx").status_code == 404
    # In use on a note: retire it, don't delete it.
    assert admin_client.delete(f"/api/settings/targetreceivers/{t['id']}").status_code == 409


def test_a_local_note_has_no_target_receiver(client, pallets):
    sync_pallets(client, pallets=[pallet("p3", barcode="B3", market_id=5, market_name="Local")])
    sync(client, note(pallets=("p3",), target_receiver_id=1, target_receiver="Someone"))
    [n] = notes(client)
    assert (n["target_receiver_id"], n["target_receiver"]) == (None, "")


def test_the_number_is_b_for_local_and_lw_for_export():
    import json as _json

    from export_forms import note_label
    from models import DeliveryNote

    def label(*markets):
        return note_label(DeliveryNote(uuid="x", note_no=1335, created_at=datetime.now(),
                                       pallets=_json.dumps([{"market_name": m} for m in markets])))
    assert (label("Local"), label("USA Sea Freight"), label("EU Air Freight"), label("Local", "EU Sea Freight"),
            label("Europe")) == ("B-1335", "LW-1335", "LW-1335", "LW-1335", "B-1335")


def test_a_note_does_not_mix_markets(client, pallets):
    from export_forms import market_group
    sync_pallets(client, pallets=[pallet("p3", barcode="B3", market_id=5, market_name="Local")])
    assert sync(client, note(pallets=("p1", "p3")))["rejected"][0]["error"] == "a note is all USA, all EU or all local pallets"
    assert sync(client, note(pallets=("p1", "p2")))["synced"] == ["n1"]
    assert [market_group(m) for m in ("USA Air Freight", "EU Sea Freight", "Local", "")] == ["usa", "eu", "local", "local"]


def test_a_receiver_copied_with_only_its_name_prints_its_market_and_city(client, admin_client, pallets):
    from io import BytesIO

    from openpyxl import load_workbook
    r = admin_client.post("/api/settings/receivers", json={"name": "ToroFruit", "market": "Joburg Mark", "city": "Johannesburg"}).json()
    sync(client, note(receiver_id=r["id"], receiver="ToroFruit"))  # its market and city set in Settings later
    c = load_workbook(BytesIO(client.get("/api/deliverynote/n1/eu_forms.xlsx").content))["Consignment"]
    assert (c["E4"].value, c["E5"].value, c["E6"].value) == ("ToroFruit", "Joburg Mark", "Johannesburg")
    assert notes(client)[0]["receiver"] == "ToroFruit"  # the note's own copy is untouched


def test_a_receivers_address_prints_between_its_market_and_city(client, admin_client, pallets):
    r = admin_client.post("/api/settings/receivers", json={"name": "ToroFruit", "market": "Joburg Mark",
                                                           "address_1": " Stand 5 ", "address_2": "", "city": "Johannesburg"}).json()
    assert (r["address_1"], r["address_2"]) == ("Stand 5", "")  # either line may stay empty
    assert admin_client.post("/api/settings/receivers", json={"name": "Bare"}).status_code == 200
    sync(client, note(receiver_id=r["id"], receiver="ToroFruit"))
    assert client.get("/api/deliverynote/n1/eu_forms.xlsx").status_code == 200  # an export note with no target receiver
    from routers.delivery_note import _live
    from db import get_packhouse_session
    s = next(client.app.dependency_overrides.get(get_packhouse_session, get_packhouse_session)())
    assert _live("n1", s).receiver == "ToroFruit\nJoburg Mark\nStand 5\nJohannesburg"
