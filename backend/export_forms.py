"""The export papers for a delivery note, as Excel or PDF, for the USA or the EU.

The USA gets three papers: DAFF's accompanying document for the phyto
inspection, the consignment note and PPECB's litchi inspection findings
sheet. The EU gets the same but no phyto document. One row per pallet line,
drawn by forms() as the farm's paper forms ("Lw 2290 #41"), one style
throughout, A4 portrait, the Phyto form on one page.
The farm's registration numbers are printed here, as the delivery note's
sender is in routers/delivery_note.py.
"""
import io
import json
import re
from datetime import date
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image
from openpyxl.drawing.fill import Blip
from openpyxl.drawing.geometry import Point2D, PositiveSize2D, PresetGeometry2D, Transform2D
from openpyxl.drawing.picture import PictureFrame
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ponytail: one exporter (AgriLink); move these to Settings -> Delivery Note when a second one ships
EXPORTER = ["AgriLink", "Ben Halliday", "0829003515"]
PRODUCER = ["Bekfontein", "Jeppes Reef rd", "Malalane"]
SENDER = ["Bekfontein Boerdery", "Jeppes Reef rd", "Malelane", "1320"]
PUC, PHC, INSPECTION_POINT, GGN = "D2087", "D 12167", "6142", "4050373828202"
PACKHOUSE_CONTACT = "0836276926"
SIZES = ("XXXL", "XXL", "XL", "L", "M", "S", "MIX")
LOGO = Path(__file__).resolve().parent / "templates" / "dalrrd_logo.jpeg"
PHYTO_ROWS = 13  # the Phyto form's ruled pallet rows, at least; more lines get more rows
FINDINGS = ["Decay", "Torn out stems", "Injuries", "Blemishes", "Sooty mould", "Incomplete sulphur", "Double litchis",
            "Absence of stems", "Long stem", "Fruit age", "Bruto mass", "Smaller 22 mm", "Uniformity", "Min. diameter"]


EXPORT_MARKET = re.compile(r"\b(usa|eu)\b", re.I)  # a pallet's market name that makes it export


def market_group(name: str) -> str:
    """A pallet's market group: usa or eu by its market name, else local. A
    delivery note holds pallets of one group."""
    hit = EXPORT_MARKET.search(name or "")
    return hit.group(1).lower() if hit else "local"


def note_label(n) -> str:
    """The note's number as printed: LW-1335 when it carries export (USA or EU)
    pallets, else B-1335 (local): one series of numbers, only the prefix differs.
    A juice note has its own series, J-0001 on."""
    export = any(EXPORT_MARKET.search(p.get("market_name") or "") for p in json.loads(n.pallets or "[]"))
    juice = json.loads(getattr(n, "bins", None) or "[]")
    if juice:
        return f"J-{n.note_no:04d}"
    return f"{'LW' if export else 'B'}-{n.note_no}"


def _count(pack: str) -> str:
    """The size count (XL) out of the pack's name (Rhino 4.5 kg XL)."""
    return next((w for w in pack.upper().replace(",", " ").split() if w in SIZES), "")


def export_rows(pallets: list) -> list:
    rows = []
    for p in pallets:
        for ln in p["lines"]:
            rows.append({"pallet_no": p["pallet_ref"], "pallet_id": p["barcode"].replace("(00)", "", 1),
                         "variety": ln.get("variety") or "", "block": (ln.get("block_name") or "").replace("Block ", ""),
                         "cartons": ln["cartons"], "date_code": p["date_code"], "count": _count(ln["packaging_name"]),
                         "pack": ln["packaging_name"]})
    return rows


def _facts(n, target: str, market: str):
    pallets = json.loads(n.pallets or "[]")
    rows = export_rows(pallets)
    return {"rows": rows, "doc_no": note_label(n), "pallets": len(pallets),
            "cartons": sum(r["cartons"] for r in rows),
            "day": date.fromisoformat(n.loaded_on).strftime("%d/%m/%Y") if n.loaded_on else "",
            "varieties": ", ".join(sorted({r["variety"] for r in rows if r["variety"]})),
            "receiver": [line for line in n.receiver.split("\n") if line.strip()][:7],
            "target": target.strip(), "market": market}


def _bytes(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# One look for every form: Arial, printed labels bold on light grey, what is
# filled in regular, every block boxed with the same thin line.
LINE = Side(style="thin", color="7F7F7F")
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
GREY = PatternFill("solid", fgColor="EDEDED")
KINDS = {  # font, fill, horizontal alignment
    "title": (Font(name="Arial", size=13, bold=True), None, "center"),
    "sub": (Font(name="Arial", size=10, bold=True), None, "center"),
    "head": (Font(name="Arial", size=9, bold=True), GREY, "center"),
    "label": (Font(name="Arial", size=9, bold=True), GREY, "left"),
    "value": (Font(name="Arial", size=10), None, "left"),
    "cell": (Font(name="Arial", size=10), None, "center"),
    "hint": (Font(name="Arial", size=8, italic=True, color="808080"), None, "center"),
}


BOX_FONT, ITEM_FONT = InlineFont(rFont="Arial", sz=16), InlineFont(rFont="Arial", sz=10)


def _boxes(items: list) -> CellRichText:
    """Tick boxes to mark by hand, each box drawn bigger than its words."""
    blocks = []
    for i, item in enumerate(items):
        blocks += [TextBlock(BOX_FONT, ("    " if i else "") + "\u2610"), TextBlock(ITEM_FONT, " " + item)]
    return CellRichText(blocks)


def _put(ws, ref, value=None, kind="value", box=True, rotate=0, top=False):
    """Write a block (a cell, or a range merged into one) in one of KINDS."""
    if ":" in ref:
        ws.merge_cells(ref)
        cells = [c for row in ws[ref] for c in row]
    else:
        cells = [ws[ref]]
    font, fill, across = KINDS[kind]
    for c in cells:  # a merged block's border and fill are drawn on every cell in it
        if box:
            c.border = BOX
        if fill:
            c.fill = fill
    first = cells[0]
    first.value = value
    first.font = font
    first.alignment = Alignment(horizontal=across, vertical="top" if top else "center", wrap_text=True,
                                text_rotation=rotate, indent=1 if across == "left" else 0)


def _sheet(wb, title, widths, heights=None):
    ws = wb.create_sheet(title)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    return ws


def _heights(ws, rows, h):
    for r in rows:
        ws.row_dimensions[r].height = h


def _phyto(wb, f):
    """DAFF's accompanying document, on one page: every pallet line in it,
    the rows still tall enough to write in, the page scaled to fit."""
    ws = _sheet(wb, "Phyto", [5, 11, 11, 9, 13, 8, 10, 10, 8, 10])
    ws.page_setup.fitToHeight = 1
    if LOGO.exists():
        logo = Image(str(LOGO))
        px = 9525  # EMU per pixel
        # Pinned at both corners and with its own size, as the farm's form has it:
        # some viewers (the Mac's preview) skip a picture without them.
        anchor = TwoCellAnchor(editAs="oneCell", _from=AnchorMarker(col=0, colOff=4 * px, row=0, rowOff=4 * px),
                               to=AnchorMarker(col=2, colOff=53 * px, row=0, rowOff=76 * px))
        anchor.pic = PictureFrame()
        anchor.pic.nvPicPr.cNvPr.name = "DALRRD logo"
        anchor.pic.blipFill.blip = Blip()
        anchor.pic.spPr.prstGeom = PresetGeometry2D(prst="rect")
        anchor.pic.spPr.xfrm = Transform2D(off=Point2D(x=4 * px, y=4 * px), ext=PositiveSize2D(cx=175 * px, cy=72 * px))
        logo.anchor = anchor
        ws.add_image(logo)
    ws.row_dimensions[1].height = 56
    _put(ws, "F1:J1", "DIRECTORATE: INSPECTION SERVICES", "sub", box=False)
    _put(ws, "A2:J2", "ACCOMPANYING DOCUMENT FOR PHYTOSANITARY INSPECTIONS FOR FRESH FRUIT", "title", box=False)
    _heights(ws, (2,), 24)

    # Left: the consignment. Right: for the PPECB inspector to fill in.
    left = [("Document no / series no", f["doc_no"]), ("Commodity", "Litchis"), ("Importing country", "USA"),
            ("Exporting company and contact details", "\n".join(EXPORTER)), ("Producer's address", "\n".join(PRODUCER)),
            ("Magisterial district", PRODUCER[-1]), ("Inspection depot", INSPECTION_POINT), ("Packhouse code", PHC),
            ("Packhouse person's contact details", PACKHOUSE_CONTACT)]
    right = [("DAFF seq. no (DAFF use only)", ""), ("Inspection date", f["day"]), ("Inspector no.", ""),
             ("Inspector signature", ""), ("Contact details and no.", "")]
    _put(ws, "F5:J5", "FOR PPECB USE ONLY", "head")
    for i, (label, value) in enumerate(left):
        r = 5 + i
        _put(ws, f"A{r}:C{r}", label, "label")
        _put(ws, f"D{r}:E{r}", value)
        ws.row_dimensions[r].height = 38 if "\n" in value else 20
    for i, (label, value) in enumerate(right):
        r = 6 + i
        _put(ws, f"F{r}:G{r}", label, "label")
        _put(ws, f"H{r}:J{r}", value)
    _put(ws, "F11:J13", "PPECB stamp", "hint", top=True)

    _put(ws, "A15:C15", "Intake document no.", "head")
    _put(ws, "D15:E15", "PUCs", "head")
    _put(ws, "F15:G15", "Total cartons", "head")
    _put(ws, "H15:J15", "Total pallets", "head")
    for ref, value in (("A16:C16", f["doc_no"]), ("D16:E16", PUC), ("F16:G16", f["cartons"]), ("H16:J16", f["pallets"])):
        _put(ws, ref, value, "cell")
    _heights(ws, (15, 16), 20)

    # The pallets, at least PHYTO_ROWS rows, the empty ones still ruled.
    head = 18
    for col, title in zip(("A", "B:C", "D", "E", "F", "G", "H", "I", "J"),
                          ("#", "Pallet ID", "PUC", "Variety", "Block no", "Carton quantity", "Date code", "Count", "Pallet nr")):
        ref = f"{col.split(':')[0]}{head}:{col.split(':')[-1]}{head}" if ":" in col else f"{col}{head}"
        _put(ws, ref, title, "head")
    ws.row_dimensions[head].height = 30
    lines = max(PHYTO_ROWS, len(f["rows"]))
    for k in range(lines):
        r = head + 1 + k
        x = f["rows"][k] if k < len(f["rows"]) else None
        values = (k + 1, x["pallet_id"], PUC, x["variety"], x["block"], x["cartons"], x["date_code"], x["count"],
                  x["pallet_no"]) if x else ("",) * 9
        for col, v in zip(("A", "B:C", "D", "E", "F", "G", "H", "I", "J"), values):
            _put(ws, f"B{r}:C{r}" if col == "B:C" else f"{col}{r}", v, "cell")
        ws.row_dimensions[r].height = 20

    # The inspectors' part.
    r = head + lines + 2
    _put(ws, f"A{r}:E{r}", "Results of inspection", "head")
    _put(ws, f"F{r}:G{r}", "Onhold", "head")
    _put(ws, f"H{r}:J{r}", "Remarks / comments by DAFF", "head")
    _put(ws, f"A{r + 1}:E{r + 4}", "DAFF stamp", "hint", top=True)
    _put(ws, f"F{r + 1}:G{r + 4}")
    _put(ws, f"H{r + 1}:J{r + 4}")
    _heights(ws, range(r + 1, r + 5), 18)
    _heights(ws, (3, 4, 14, 17, r - 1), 6)  # the gaps between blocks: the page is fitted to its height
    for i, (a, b) in enumerate((("TempTale", "Pallet nr"), ("Signature of DAFF inspector", "Signature of foreign inspector"))):
        q = r + 5 + i
        _put(ws, f"A{q}:B{q}", a, "label")
        _put(ws, f"C{q}:E{q}")
        _put(ws, f"F{q}:G{q}", b, "label")
        _put(ws, f"H{q}:J{q}")
        ws.row_dimensions[q].height = 28


def _consignment(wb, f):
    ws = _sheet(wb, "Consignment", [4, 8, 19, 8, 9, 11, 8, 7, 14, 15, 8])
    _put(ws, "A1:K1", "CONSIGNMENT NOTE", "title", box=False)
    ws.row_dimensions[1].height = 26
    _put(ws, "A3:D3", "Sender", "head")
    _put(ws, "E3:H3", "Receiver", "head")
    _put(ws, "I3:K3", "Consignment", "head")
    sender = [(line, None) for line in SENDER] + [("PUC code", PUC), ("PHC code", PHC), ("GGN", GGN),
                                                 ("GlobalG.A.P. Certified", None)]
    codes = [("Consignment nr", f["doc_no"]), ("Date", f["day"]), ("Inspection point", INSPECTION_POINT),
             ("Target market", f["market"])]
    for i, (a, b) in enumerate(sender):
        r = 4 + i
        if b is None:
            _put(ws, f"A{r}:D{r}", a)
        else:
            _put(ws, f"A{r}:B{r}", a, "label")
            _put(ws, f"C{r}:D{r}", b)
        _put(ws, f"E{r}:H{r}", f["receiver"][i] if i < len(f["receiver"]) else "")
        if i < len(codes):
            _put(ws, f"I{r}", codes[i][0], "label")
            _put(ws, f"J{r}:K{r}", codes[i][1])
        ws.row_dimensions[r].height = 18
    _put(ws, "I8:I10", "Target receiver", "label", top=True)
    _put(ws, "J8:K10", f["target"], top=True)

    head = 13
    titles = ("#", "Pallet nr", "Pallet ID", "Start date", "Commodity", "Variety", "Grade", "Count", "Block nr",
              "Package type", "Units")
    for i, title in enumerate(titles):
        _put(ws, f"{chr(65 + i)}{head}", title, "head")
    ws.row_dimensions[head].height = 28
    ws.print_title_rows = f"{head}:{head}"
    for k, x in enumerate(f["rows"]):
        r = head + 1 + k
        for i, v in enumerate((k + 1, x["pallet_no"], x["pallet_id"], x["date_code"], "Litchis", x["variety"], "Class 1",
                               x["count"], x["block"], x["pack"], x["cartons"])):
            _put(ws, f"{chr(65 + i)}{r}", v, "cell")
        ws.row_dimensions[r].height = 18
    r = head + len(f["rows"]) + 1
    _put(ws, f"I{r}:J{r}", "Total units", "label")
    _put(ws, f"K{r}", f["cartons"], "cell")
    ws.row_dimensions[r].height = 20

    # For the PPECB inspector, the TempTale beside it; then who signs for the load.
    r += 2
    for i, label in enumerate(("Inspection date", "Inspector nr", "Contact nr", "Signature")):
        _put(ws, f"A{r + i}:C{r + i}", label, "label")
        _put(ws, f"D{r + i}:F{r + i}", f["day"] if i == 0 else "")
        ws.row_dimensions[r + i].height = 22
    _put(ws, f"H{r}:I{r}", "TempTale", "label")
    _put(ws, f"J{r}:K{r}")
    _put(ws, f"H{r + 1}:I{r + 1}", "Pallet nr", "label")
    _put(ws, f"J{r + 1}:K{r + 1}")
    r += 5
    for i, label in enumerate(("Driver signature", "Receiving signature", "Packhouse representative")):
        _put(ws, f"A{r + i}:C{r + i}", label, "label")
        _put(ws, f"D{r + i}:F{r + i}")
        _put(ws, f"G{r + i}:H{r + i}", "Date", "label")
        _put(ws, f"I{r + i}:K{r + i}")
        ws.row_dimensions[r + i].height = 30


def _inspection(wb, f):
    """PPECB's inspection findings sheet: the findings are for the inspector."""
    ws = _sheet(wb, "Inspection Sheet", [21, 8, 6, 8] + [4.6] * len(FINDINGS) + [9])
    last = chr(65 + 4 + len(FINDINGS))  # S
    _put(ws, "A1:N1", "PERISHABLE PRODUCTS EXPORT CONTROL BOARD", "title", box=False)
    _put(ws, f"O1:{last}1", "PP04.01-05F08 L  Rev: 0", "hint", box=False)
    _put(ws, "A2:N2", "INSPECTION FINDINGS SHEET - LITCHIS", "sub", box=False)
    _put(ws, f"O2:{last}2", "PPECB FINDINGS", "sub", box=False)
    _heights(ws, (1, 2), 22)
    _put(ws, "A4", "Mark the option (√)", "label")
    for ref, option in (("B4:E4", "Min standard"), ("F4:J4", "Template"), ("K4:O4", "Special market"), (f"P4:{last}4", "Dispensation")):
        _put(ws, ref, _boxes([option]), "cell")
    ws.row_dimensions[4].height = 24
    pairs = [[("Inspector and no", ""), ("Exporter", EXPORTER[0]), ("PHC", PHC), ("PUC", PUC)],
             [("Signature", ""), ("Insp. point", INSPECTION_POINT), ("Date", f["day"]), ("Product", "Litchis")],
             [("Consignment note nr", f["doc_no"]), ("Cultivar", f["varieties"]), ("", ""), ("", "")]]
    spans = (("A", "B:E"), ("F:G", "H:J"), ("K:L", "M:O"), ("P:Q", f"R:{last}"))
    for i, row in enumerate(pairs):
        r = 5 + i
        for (lab, val), (lc, vc) in zip(row, spans):
            lref = f"{lc.split(':')[0]}{r}:{lc.split(':')[-1]}{r}" if ":" in lc else f"{lc}{r}"
            _put(ws, lref, lab, "label")
            _put(ws, f"{vc.split(':')[0]}{r}:{vc.split(':')[1]}{r}", val)
        ws.row_dimensions[r].height = 22

    head = 9
    for i, title in enumerate(["Pallet no / pallet barcode ID", "Net mass", "Class", "Target market"] + FINDINGS
                              + ["C - comply\nR - rejected"]):
        _put(ws, f"{chr(65 + i)}{head}", title, "head", rotate=90 if 4 <= i < 4 + len(FINDINGS) else 0)
    ws.row_dimensions[head].height = 96
    ws.print_title_rows = f"{head}:{head}"
    for k, x in enumerate(f["rows"]):
        r = head + 1 + k
        for i in range(5 + len(FINDINGS)):
            _put(ws, f"{chr(65 + i)}{r}", {0: x["pallet_id"], 2: 1, 3: f["market"]}.get(i, ""), "cell")
        ws.row_dimensions[r].height = 20

    r = head + len(f["rows"]) + 2
    checks = [("Marking requirements", "Product, PUC, Class, Size diameter, Count, Origin, Address, Cultivar, Date code, Treated with ..."),
              ("Containers", "Clean, Dry, Undamaged, Suitable, No odour, No arthropoda, No fungus"),
              ("Pallet requirements", "Clean, Fungi, Bark, Undamaged, Same dimensions on the same layers, Arthropoda, Containers square, Secured, Covered, ISPM requirements"),
              ("Traceability", "PUC, Date code, Pallet ID, Address, Consignment note/s"),
              ("Packing requirements", "Pre-pack, Same cultivar, Same size, Same quality, Same ripeness")]
    for i, (label, items) in enumerate(checks):
        _put(ws, f"A{r + i}", label, "label")
        _put(ws, f"B{r + i}:{last}{r + i}", _boxes(items.split(", ")))
        # ponytail: a character count for the wrap (about 95 to a line across B:S, a box and its gap
        # about 4), not a measured width
        ws.row_dimensions[r + i].height = 22 * -(-sum(len(x) + 4 for x in items.split(", ")) // 95) + 4
    r += len(checks)
    _put(ws, f"A{r}", "Passed for export stickers", "label")
    _put(ws, f"B{r}:C{r}", "From", "label")
    _put(ws, f"D{r}:J{r}")
    _put(ws, f"K{r}:L{r}", "To", "label")
    _put(ws, f"M{r}:{last}{r}")
    _put(ws, f"A{r + 1}", "Remarks", "label", top=True)
    _put(ws, f"B{r + 1}:{last}{r + 1}")
    ws.row_dimensions[r].height = 22
    ws.row_dimensions[r + 1].height = 60


def _workbook(n, target: str, market: str) -> Workbook:
    f = _facts(n, target, market)
    wb = Workbook()
    wb.remove(wb.active)
    if market == "USA":  # the EU has no phyto document
        _phyto(wb, f)
    _consignment(wb, f)
    _inspection(wb, f)
    return wb


def forms(n, target: str = "", market: str = "USA") -> bytes:
    """The papers drawn as the farm's forms; target is the note's target receiver."""
    return _bytes(_workbook(n, target, market))


def forms_pdf(n, target: str = "", market: str = "USA") -> bytes:
    """The same papers as one PDF, every sheet drawn from the workbook forms()
    makes, so the two never differ: each sheet fitted to the A4 width (the
    Phyto form to its height too), a long one carried on to the next page
    with its header row again."""
    pdf = FPDF(unit="pt", format="A4")
    pdf.set_auto_page_break(False)
    pdf.c_margin = 0  # the cells are padded by _text
    for ws in _workbook(n, target, market).worksheets:
        _sheet_pdf(pdf, ws)
    return bytes(pdf.output())


# Drawing a worksheet in fpdf2, for the few things _put and _sheet use.
EMU_PT = 12700


def _latin(v) -> str:
    # fpdf2's built-in fonts are latin-1: the tick is a v, a tick box is drawn (_rich).
    return ("" if v is None else str(v)).replace("\u221a", "v").encode("latin-1", "replace").decode("latin-1")


def _rgb(color, default=(0, 0, 0)):
    rgb = getattr(color, "rgb", None)
    return tuple(int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4)) if isinstance(rgb, str) else default


def _font(pdf, font, scale, size=None, bold=None):
    style = ("B" if (font.b if bold is None else bold) else "") + ("I" if font.i else "")
    pdf.set_font("Helvetica", style, (size or font.sz or 11) * scale)
    pdf.set_text_color(*_rgb(font.color))


def _rich(pdf, cell, x, y, w, h, scale):
    """Tick boxes (_boxes) laid out word by word: each box a square drawn
    with its item beside it, an item kept whole on a line."""
    _font(pdf, cell.font, scale, size=10)
    pad, lh = 4 * scale, 12 * scale
    size = pdf.font_size * 0.9
    items = [b.text.strip() for b in cell.value if b.text.strip() != "\u2610"]
    widths = [size + 3 * scale + pdf.get_string_width(_latin(t)) for t in items]
    lines, line = [], []
    for t, iw in zip(items, widths):
        if line and sum(x[1] for x in line) + 8 * scale * len(line) + iw > w - 2 * pad:
            lines.append(line)
            line = []
        line.append((t, iw))
    lines.append(line)
    centre = cell.alignment.horizontal == "center"
    ty = y + (h - lh * len(lines)) / 2
    for line in lines:
        used = sum(x[1] for x in line) + 8 * scale * (len(line) - 1)
        tx = x + ((w - used) / 2 if centre else pad)
        for t, iw in line:
            pdf.rect(tx, ty + (lh - size) / 2, size, size)
            pdf.set_xy(tx + size + 2 * scale, ty)
            pdf.cell(iw - size, lh, _latin(t))
            tx += iw + 8 * scale
        ty += lh


def _text(pdf, cell, x, y, w, h, scale):
    a = cell.alignment
    _font(pdf, cell.font, scale)
    lh = pdf.font_size * 1.2
    if a.text_rotation == 90:  # read bottom to top: drawn across a box turned on its side
        cx, cy = x + w / 2, y + h / 2
        x, y, w, h = cx - h / 2, cy - w / 2, h, w
    pad = 5 * scale if a.indent else 1.5 * scale
    text = _latin(cell.value)
    widest = max((pdf.get_string_width(word) for word in text.split()), default=0)
    if widest > w - 2 * pad:  # a long number (a pallet ID) is shrunk to fit, never broken
        pdf.set_font_size(pdf.font_size_pt * (w - 2 * pad) / widest)
        lh = pdf.font_size * 1.2
    lines = len(pdf.multi_cell(w - 2 * pad, lh, text, dry_run=True, output="LINES"))
    top = y + 2 * scale if a.vertical == "top" else y + (h - lines * lh) / 2
    align = {"center": "C", "right": "R"}.get(a.horizontal, "L")
    if a.text_rotation == 90:
        with pdf.rotation(90, cx, cy):
            pdf.set_xy(x + pad, top)
            pdf.multi_cell(w - 2 * pad, lh, text, align=align)
    else:
        pdf.set_xy(x + pad, top)
        pdf.multi_cell(w - 2 * pad, lh, text, align=align)


def _sheet_pdf(pdf, ws):
    cols = [ws.column_dimensions[chr(65 + i)].width or 8.43 for i in range(ws.max_column)]
    colw = [(w * 7 + 5) * 0.75 for w in cols]  # Excel's character width in points
    rowh = {r: ws.row_dimensions[r].height or 15 for r in range(1, ws.max_row + 1)}
    m = ws.page_margins
    left, top = m.left * 72, m.top * 72
    room_w, room_h = pdf.w - 2 * left, pdf.h - 2 * top
    scale = min(1, room_w / sum(colw))
    if ws.page_setup.fitToHeight == 1:
        scale = min(scale, room_h / sum(rowh.values()))
    x0 = (pdf.w - sum(colw) * scale) / 2
    xs = [x0 + sum(colw[:i]) * scale for i in range(len(colw) + 1)]
    spans = {}  # a merged block's top-left cell: its last row and column
    hidden = set()
    for rng in ws.merged_cells.ranges:
        spans[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
        hidden |= {(r, c) for r in range(rng.min_row, rng.max_row + 1) for c in range(rng.min_col, rng.max_col + 1)}
    titles = ws.print_title_rows
    head = [int(r) for r in titles.replace("$", "").split(":")] if titles else []
    pdf.set_draw_color(*_rgb(LINE.color))
    pdf.set_line_width(0.5 * scale)

    def draw_row(r, y):
        for c in range(1, len(colw) + 1):
            if (r, c) in hidden and (r, c) not in spans:
                continue
            r2, c2 = spans.get((r, c), (r, c))
            cell = ws.cell(r, c)
            x, w = xs[c - 1], xs[c2] - xs[c - 1]
            h = sum(rowh[k] for k in range(r, r2 + 1)) * scale
            if cell.fill.fill_type == "solid":
                pdf.set_fill_color(*_rgb(cell.fill.fgColor, (255, 255, 255)))
                pdf.rect(x, y, w, h, style="F")
            if cell.border.left.style:
                pdf.rect(x, y, w, h)
            if isinstance(cell.value, CellRichText):
                _rich(pdf, cell, x, y, w, h, scale)
            elif cell.value not in (None, ""):
                _text(pdf, cell, x, y, w, h, scale)

    pdf.add_page()
    for img in ws._images:
        a = img.anchor
        ext = a.pic.spPr.xfrm.ext
        x = xs[a._from.col] + a._from.colOff / EMU_PT * scale
        y = top + sum(rowh[k] for k in range(1, a._from.row + 1)) * scale + a._from.rowOff / EMU_PT * scale
        pdf.image(str(LOGO), x, y, ext.cx / EMU_PT * scale, ext.cy / EMU_PT * scale)
    y = top
    for r in range(1, ws.max_row + 1):
        # A merged block is kept whole: the page breaks before its first row.
        tall = max([sum(rowh[k] for k in range(r, spans[(r, c)][0] + 1)) for c in range(1, len(colw) + 1)
                    if (r, c) in spans] + [rowh[r]]) * scale
        if y + tall > top + room_h + 0.5:
            pdf.add_page()
            y = top
            for t in range(head[0], head[-1] + 1) if head and r > head[-1] else ():
                draw_row(t, y)
                y += rowh[t] * scale
        draw_row(r, y)
        y += rowh[r] * scale

